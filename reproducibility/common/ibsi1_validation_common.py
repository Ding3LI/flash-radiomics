from __future__ import annotations

import math
import time
from dataclasses import dataclass
import csv
import json
from pathlib import Path
from typing import Callable

import numpy as np
import SimpleITK as sitk
from openpyxl import load_workbook

from reproducibility.common.official_benchmark_common import relative_to_workspace, write_csv, write_json


WORKSPACE_ROOT = Path(__file__).resolve().parents[2]

DEFAULT_IBSI1_DATA_ROOT = WORKSPACE_ROOT / "reproducibility" / "data" / "ibsi" / "data_sets"
DEFAULT_IBSI1_TEMPLATE = WORKSPACE_ROOT / "reproducibility" / "data" / "ibsi" / "IBSI-1-submission-table.xlsx"
DEFAULT_IBSI1_CONFIG_JSON = WORKSPACE_ROOT / "reproducibility" / "data" / "ibsi" / "ibsi1_configs_A_to_E.json"
DEFAULT_IBSI_HIGH_CONSENSUS_MAP = WORKSPACE_ROOT / "reproducibility" / "ibsi" / "ibsi_high_consensus_165_feature_map.csv"

IBSI1_CASE_ID = "ibsi1_ct_phantom"
IBSI1_LABEL = 1
IBSI1_ROI_NAME = "GTV-1"

ALL_ROWS_STATUS_ORDER = ("pass", "fail", "unsupported", "not_standardized", "error")

TEXTURE_SUFFIX_BY_MODE = {
    "2D": {
        "glcm": "2_5D_avg",
        "glrlm": "2_5D_avg",
        "glszm": "2_5D",
        "gldzm": "2_5D",
        "gldm": "2_5D",
        "ngtdm": "2_5D",
    },
    "3D": {
        "glcm": "3D_avg",
        "glrlm": "3D_avg",
        "glszm": "3D",
        "gldzm": "3D",
        "gldm": "3D",
        "ngtdm": "3D",
    },
}

GLCM_TAGS = {
    "Autocorrelation": "cm_auto_corr",
    "ClusterProminence": "cm_clust_prom",
    "ClusterShade": "cm_clust_shade",
    "ClusterTendency": "cm_clust_tend",
    "Contrast": "cm_contrast",
    "Correlation": "cm_corr",
    "DifferenceAverage": "cm_diff_avg",
    "DifferenceEntropy": "cm_diff_entr",
    "DifferenceVariance": "cm_diff_var",
    "Dissimilarity": "cm_dissimilarity",
    "Id": "cm_inv_diff",
    "Idm": "cm_inv_diff_mom",
    "Idmn": "cm_inv_diff_mom_norm",
    "Idn": "cm_inv_diff_norm",
    "Imc1": "cm_info_corr1",
    "Imc2": "cm_info_corr2",
    "InverseVariance": "cm_inv_var",
    "JointAverage": "cm_joint_avg",
    "JointEnergy": "cm_energy",
    "JointEntropy": "cm_joint_entr",
    "MaximumProbability": "cm_joint_max",
    "SumAverage": "cm_sum_avg",
    "SumEntropy": "cm_sum_entr",
    "SumVariance": "cm_sum_var",
    "SumSquares": "cm_joint_var",
}

GLRLM_TAGS = {
    "GrayLevelNonUniformity": "rlm_glnu",
    "GrayLevelNonUniformityNormalized": "rlm_glnu_norm",
    "GrayLevelVariance": "rlm_gl_var",
    "HighGrayLevelRunEmphasis": "rlm_hgre",
    "LongRunEmphasis": "rlm_lre",
    "LongRunHighGrayLevelEmphasis": "rlm_lrhge",
    "LongRunLowGrayLevelEmphasis": "rlm_lrlge",
    "LowGrayLevelRunEmphasis": "rlm_lgre",
    "RunEntropy": "rlm_rl_entr",
    "RunLengthNonUniformity": "rlm_rlnu",
    "RunLengthNonUniformityNormalized": "rlm_rlnu_norm",
    "RunPercentage": "rlm_r_perc",
    "RunVariance": "rlm_rl_var",
    "ShortRunEmphasis": "rlm_sre",
    "ShortRunHighGrayLevelEmphasis": "rlm_srhge",
    "ShortRunLowGrayLevelEmphasis": "rlm_srlge",
}

GLSZM_TAGS = {
    "GrayLevelNonUniformity": "szm_glnu",
    "GrayLevelNonUniformityNormalized": "szm_glnu_norm",
    "GrayLevelVariance": "szm_gl_var",
    "HighGrayLevelZoneEmphasis": "szm_hgze",
    "LargeAreaEmphasis": "szm_lze",
    "LargeAreaHighGrayLevelEmphasis": "szm_lzhge",
    "LargeAreaLowGrayLevelEmphasis": "szm_lzlge",
    "LowGrayLevelZoneEmphasis": "szm_lgze",
    "SizeZoneNonUniformity": "szm_zsnu",
    "SizeZoneNonUniformityNormalized": "szm_zsnu_norm",
    "SmallAreaEmphasis": "szm_sze",
    "SmallAreaHighGrayLevelEmphasis": "szm_szhge",
    "SmallAreaLowGrayLevelEmphasis": "szm_szlge",
    "ZoneEntropy": "szm_zs_entr",
    "ZonePercentage": "szm_z_perc",
    "ZoneVariance": "szm_zs_var",
}

GLDZM_TAGS = {
    "GrayLevelNonUniformity": "dzm_glnu",
    "GrayLevelNonUniformityNormalized": "dzm_glnu_norm",
    "GrayLevelVariance": "dzm_gl_var",
    "HighGrayLevelZoneEmphasis": "dzm_hgze",
    "LargeDistanceEmphasis": "dzm_lde",
    "LargeDistanceHighGrayLevelEmphasis": "dzm_ldhge",
    "LargeDistanceLowGrayLevelEmphasis": "dzm_ldlge",
    "LowGrayLevelZoneEmphasis": "dzm_lgze",
    "SmallDistanceEmphasis": "dzm_sde",
    "SmallDistanceHighGrayLevelEmphasis": "dzm_sdhge",
    "SmallDistanceLowGrayLevelEmphasis": "dzm_sdlge",
    "ZoneDistanceEntropy": "dzm_zd_entr",
    "ZoneDistanceNonUniformity": "dzm_zdnu",
    "ZoneDistanceNonUniformityNormalized": "dzm_zdnu_norm",
    "ZoneDistanceVariance": "dzm_zd_var",
    "ZonePercentage": "dzm_z_perc",
}

GLDM_TAGS = {
    "DependenceEntropy": "ngl_dc_entr",
    "DependenceNonUniformity": "ngl_dcnu",
    "DependenceNonUniformityNormalized": "ngl_dcnu_norm",
    "DependenceVariance": "ngl_dc_var",
    "GrayLevelNonUniformity": "ngl_glnu",
    "GrayLevelVariance": "ngl_gl_var",
    "HighGrayLevelEmphasis": "ngl_hgce",
    "LargeDependenceEmphasis": "ngl_hde",
    "LargeDependenceHighGrayLevelEmphasis": "ngl_hdhge",
    "LargeDependenceLowGrayLevelEmphasis": "ngl_hdlge",
    "LowGrayLevelEmphasis": "ngl_lgce",
    "SmallDependenceEmphasis": "ngl_lde",
    "SmallDependenceHighGrayLevelEmphasis": "ngl_ldhge",
    "SmallDependenceLowGrayLevelEmphasis": "ngl_ldlge",
}

NGTDM_TAGS = {
    "Busyness": "ngt_busyness",
    "Coarseness": "ngt_coarseness",
    "Complexity": "ngt_complexity",
    "Contrast": "ngt_contrast",
    "Strength": "ngt_strength",
}

IH_TAGS = {
    "Mean": "ih_mean",
    "Variance": "ih_var",
    "Skewness": "ih_skew",
    "Kurtosis": "ih_kurt",
    "Median": "ih_median",
    "Minimum": "ih_min",
    "10Percentile": "ih_p10",
    "90Percentile": "ih_p90",
    "Maximum": "ih_max",
    "Mode": "ih_mode",
    "InterquartileRange": "ih_iqr",
    "Range": "ih_range",
    "MeanAbsoluteDeviation": "ih_mad",
    "RobustMeanAbsoluteDeviation": "ih_rmad",
    "MedianAbsoluteDeviation": "ih_medad",
    "CoefficientOfVariation": "ih_cov",
    "QuartileCoefficientOfDispersion": "ih_qcod",
    "Entropy": "ih_entropy",
    "Uniformity": "ih_uniformity",
    "MaximumHistogramGradient": "ih_max_grad",
    "MaximumHistogramGradientIntensity": "ih_max_grad_g",
    "MinimumHistogramGradient": "ih_min_grad",
    "MinimumHistogramGradientIntensity": "ih_min_grad_g",
}

IVH_TAGS = {
    "VolumeFractionAt10Intensity": "ivh_v10",
    "VolumeFractionAt90Intensity": "ivh_v90",
    "IntensityAt10Volume": "ivh_i10",
    "IntensityAt90Volume": "ivh_i90",
    "VolumeFractionDifferenceBetween10And90Intensity": "ivh_diff_v10_v90",
    "IntensityDifferenceBetween10And90Volume": "ivh_diff_i10_i90",
}

FIRSTORDER_TAGS = {
    "10Percentile": "stat_p10",
    "90Percentile": "stat_p90",
    "Energy": "stat_energy",
    "Entropy": "ih_entropy",
    "InterquartileRange": "stat_iqr",
    "Kurtosis": "stat_kurt",
    "Maximum": "stat_max",
    "Mean": "stat_mean",
    "MeanAbsoluteDeviation": "stat_mad",
    "Median": "stat_median",
    "Minimum": "stat_min",
    "Range": "stat_range",
    "RobustMeanAbsoluteDeviation": "stat_rmad",
    "RootMeanSquared": "stat_rms",
    "Skewness": "stat_skew",
    "Uniformity": "ih_uniformity",
    "Variance": "stat_var",
}

SHAPE_RESULT_TAGS = {
    "Elongation": "morph_pca_elongation",
    "Flatness": "morph_pca_flatness",
    "LeastAxisLength": "morph_pca_least_axis",
    "MajorAxisLength": "morph_pca_maj_axis",
    "Maximum3DDiameter": "morph_diam",
    "MeshVolume": "morph_volume",
    "MinorAxisLength": "morph_pca_min_axis",
    "Sphericity": "morph_sphericity",
    "SurfaceArea": "morph_area_mesh",
    "SurfaceVolumeRatio": "morph_av",
    "VoxelVolume": "morph_vol_approx",
}

# Additional feature tags for IBSI validation reporting.
VALIDATION_FIRSTORDER_EXTRA_TAGS = {
    "CoefficientOfVariation": "stat_cov",
    "MedianAbsoluteDeviation": "stat_medad",
    "QuartileCoefficientOfDispersion": "stat_qcod",
}

VALIDATION_GLDM_EXTRA_TAGS = {
    "DependenceCountEnergy": "ngl_dc_energy",
    "DependenceCountPercentage": "ngl_dc_perc",
    "GrayLevelNonUniformityNormalized": "ngl_glnu_norm",
}

VALIDATION_SHAPE_EXTRA_TAGS = {
    "Asphericity": "morph_asphericity",
    "CentreOfMassShift": "morph_com",
    "IntegratedIntensity": "morph_integ_int",
    "VolumeDensityAABB": "morph_vol_dens_aabb",
    "AreaDensityAABB": "morph_area_dens_aabb",
    "VolumeDensityAEE": "morph_vol_dens_aee",
    "AreaDensityAEE": "morph_area_dens_aee",
    "VolumeDensityConvexHull": "morph_vol_dens_conv_hull",
    "AreaDensityConvexHull": "morph_area_dens_conv_hull",
}


@dataclass
class PreparedConfigCase:
    config_id: str
    extraction_mode: str
    note: str
    initial_image: sitk.Image
    initial_mask: sitk.Image
    interpolated_image: sitk.Image
    interpolated_mask: sitk.Image
    morph_mask: sitk.Image
    intensity_mask: sitk.Image
    discretization: dict
    force2d: bool
    force2d_dimension: int
    texture_settings: dict


def _copy_image_info(src: sitk.Image, arr: np.ndarray, pixel_type: int) -> sitk.Image:
    image = sitk.GetImageFromArray(arr)
    image = sitk.Cast(image, pixel_type)
    image.CopyInformation(src)
    return image


def _round_image_to_intensity_grid(image: sitk.Image) -> sitk.Image:
    arr = sitk.GetArrayFromImage(image).astype(np.float32, copy=False)
    rounded = np.rint(arr).astype(np.float32, copy=False)
    return _copy_image_info(image, rounded, sitk.sitkFloat32)


def _read_rounded_phantom(image_path: Path, mask_path: Path) -> tuple[sitk.Image, sitk.Image]:
    image = _round_image_to_intensity_grid(sitk.Cast(sitk.ReadImage(str(image_path)), sitk.sitkFloat32))
    mask_arr = sitk.GetArrayFromImage(sitk.ReadImage(str(mask_path))).astype(np.float32, copy=False)
    mask_bin = (mask_arr > 0.5).astype(np.uint8, copy=False)
    mask = _copy_image_info(image, mask_bin, sitk.sitkUInt8)
    return image, mask


def _sitk_interpolator(method: str | None) -> int:
    method_text = str(method or "").strip().lower()
    if method_text in {"bilinear", "trilinear", "linear"}:
        return sitk.sitkLinear
    if method_text in {"tricubic_spline", "bspline", "cubic"}:
        return sitk.sitkBSpline
    return sitk.sitkLinear


def _compute_resampled_size(size: tuple[int, ...], spacing: tuple[float, ...], target_spacing: tuple[float, ...]) -> list[int]:
    result: list[int] = []
    for dim_size, dim_spacing, dim_target in zip(size, spacing, target_spacing):
        physical_extent = float(dim_size) * float(dim_spacing)
        new_size = int(math.ceil(physical_extent / float(dim_target)))
        result.append(max(1, new_size))
    return result


def _resample_image(image: sitk.Image, spacing: tuple[float, ...], interpolator: int, pixel_type: int) -> sitk.Image:
    from scipy.ndimage import map_coordinates

    arr = sitk.GetArrayFromImage(image).astype(np.float64, copy=False)
    original_size_xyz = tuple(int(v) for v in image.GetSize())
    original_spacing_xyz = tuple(float(v) for v in image.GetSpacing())
    new_size_xyz = _compute_resampled_size(original_size_xyz, original_spacing_xyz, spacing)

    original_dim_zyx = np.array(arr.shape, dtype=np.float64)
    original_spacing_zyx = np.array(original_spacing_xyz[::-1], dtype=np.float64)
    new_spacing_zyx = np.array(tuple(float(v) for v in spacing)[::-1], dtype=np.float64)
    new_dim_zyx = np.array(new_size_xyz[::-1], dtype=np.int64)
    grid_spacing = new_spacing_zyx / original_spacing_zyx
    grid_origin = 0.5 * (original_dim_zyx - 1.0) - 0.5 * (new_dim_zyx - 1.0) * grid_spacing

    z_map, y_map, x_map = np.mgrid[: new_dim_zyx[0], : new_dim_zyx[1], : new_dim_zyx[2]]
    coordinates = np.array(
        [
            z_map.reshape(-1) * grid_spacing[0] + grid_origin[0],
            y_map.reshape(-1) * grid_spacing[1] + grid_origin[1],
            x_map.reshape(-1) * grid_spacing[2] + grid_origin[2],
        ],
        dtype=np.float64,
    )
    order = 1 if interpolator == sitk.sitkLinear else 3
    sampled = map_coordinates(arr, coordinates, order=order, mode="nearest")
    sampled = sampled.reshape(tuple(int(v) for v in new_dim_zyx))

    out = sitk.GetImageFromArray(sampled)
    out = sitk.Cast(out, pixel_type)
    out.SetSpacing(tuple(float(v) for v in spacing))
    out.SetDirection(image.GetDirection())
    out.SetOrigin(image.GetOrigin())
    return out


def _resample_case(image: sitk.Image, mask: sitk.Image, config: dict) -> tuple[sitk.Image, sitk.Image]:
    interpolation = dict(config.get("interpolation", {}))
    if not interpolation.get("enabled"):
        return image, mask

    spacing_cfg = interpolation.get("resampled_voxel_spacing_mm")
    if not isinstance(spacing_cfg, list):
        return image, mask

    original_spacing = tuple(float(v) for v in image.GetSpacing())
    if interpolation.get("spacing_scope") == "axial":
        target_spacing = (float(spacing_cfg[0]), float(spacing_cfg[1]), original_spacing[2])
    else:
        target_spacing = tuple(float(v) for v in spacing_cfg)

    image_interp = _sitk_interpolator(interpolation.get("image_interpolation_method"))
    mask_interp = _sitk_interpolator(interpolation.get("roi_interpolation_method"))

    resampled_image = _resample_image(image, target_spacing, image_interp, sitk.sitkFloat32)
    resampled_image = _round_image_to_intensity_grid(resampled_image)

    float_mask = sitk.Cast(mask, sitk.sitkFloat32)
    resampled_mask_float = _resample_image(float_mask, target_spacing, mask_interp, sitk.sitkFloat32)
    threshold = float(interpolation.get("roi_partial_mask_volume", 0.5) or 0.5)
    mask_arr = sitk.GetArrayFromImage(resampled_mask_float).astype(np.float32, copy=False)
    mask_bin = (mask_arr >= threshold).astype(np.uint8, copy=False)
    resampled_mask = _copy_image_info(resampled_image, mask_bin, sitk.sitkUInt8)
    return resampled_image, resampled_mask


def _apply_resegmentation(image: sitk.Image, mask: sitk.Image, config: dict) -> sitk.Image:
    image_arr = sitk.GetArrayFromImage(image).astype(np.float64, copy=False)
    mask_arr = sitk.GetArrayFromImage(mask).astype(bool, copy=False)
    intensity_mask = np.array(mask_arr, copy=True)

    resegmentation = dict(config.get("resegmentation", {}))
    range_hu = resegmentation.get("range_hu")
    if isinstance(range_hu, list) and len(range_hu) == 2:
        lo, hi = sorted(float(v) for v in range_hu)
        intensity_mask &= image_arr >= lo
        intensity_mask &= image_arr <= hi

    outlier = resegmentation.get("outlier_filtering")
    if isinstance(outlier, dict) and str(outlier.get("type", "")).lower() == "sigma" and np.any(intensity_mask):
        sigma = float(outlier.get("value", 0.0))
        roi_values = image_arr[intensity_mask]
        mean_value = float(np.mean(roi_values))
        std_value = float(np.std(roi_values))
        lo = mean_value - sigma * std_value
        hi = mean_value + sigma * std_value
        intensity_mask &= image_arr >= lo
        intensity_mask &= image_arr <= hi

    return _copy_image_info(image, intensity_mask.astype(np.uint8, copy=False), sitk.sitkUInt8)


def _percentile(values: np.ndarray, q: float) -> float:
    return float(np.percentile(values, q))


def _robust_values(values: np.ndarray) -> np.ndarray:
    if values.size == 0:
        return values
    p10 = _percentile(values, 10.0)
    p90 = _percentile(values, 90.0)
    robust = values[(values >= p10) & (values <= p90)]
    return robust if robust.size > 0 else values


def _safe_skew(values: np.ndarray) -> float:
    std_value = float(np.std(values))
    if std_value <= 0.0:
        return 0.0
    centered = values - float(np.mean(values))
    return float(np.mean(centered**3) / (std_value**3))


def _safe_kurtosis(values: np.ndarray) -> float:
    std_value = float(np.std(values))
    if std_value <= 0.0:
        return 0.0
    centered = values - float(np.mean(values))
    return float(np.mean(centered**4) / (std_value**4) - 3.0)


def _mode_from_values(values: np.ndarray) -> float:
    uniq, counts = np.unique(values, return_counts=True)
    if uniq.size == 0:
        return float("nan")
    return float(uniq[int(np.argmax(counts))])


def _histogram_gradient_features(values: np.ndarray) -> dict[str, float]:
    uniq, counts = np.unique(values.astype(np.int64, copy=False), return_counts=True)
    if uniq.size <= 1:
        value = float(uniq[0]) if uniq.size == 1 else float("nan")
        return {
            "ih_max_grad": 0.0,
            "ih_max_grad_g": value,
            "ih_min_grad": 0.0,
            "ih_min_grad_g": value,
        }

    gradients = np.diff(counts.astype(np.float64, copy=False))
    max_idx = int(np.argmax(gradients))
    min_idx = int(np.argmin(gradients))
    return {
        "ih_max_grad": float(gradients[max_idx]),
        "ih_max_grad_g": float(uniq[max_idx + 1]),
        "ih_min_grad": float(gradients[min_idx]),
        "ih_min_grad_g": float(uniq[min_idx + 1]),
    }


def _basic_distribution_features(values: np.ndarray, prefix: str) -> dict[str, float]:
    if values.size == 0:
        return {}

    mean_value = float(np.mean(values))
    std_value = float(np.std(values))
    median_value = float(np.median(values))
    p25 = _percentile(values, 25.0)
    p75 = _percentile(values, 75.0)
    robust = _robust_values(values)
    robust_mean = float(np.mean(robust))

    features = {
        f"{prefix}_mean": mean_value,
        f"{prefix}_var": float(np.var(values)),
        f"{prefix}_skew": _safe_skew(values),
        f"{prefix}_kurt": _safe_kurtosis(values),
        f"{prefix}_median": median_value,
        f"{prefix}_min": float(np.min(values)),
        f"{prefix}_p10": _percentile(values, 10.0),
        f"{prefix}_p90": _percentile(values, 90.0),
        f"{prefix}_max": float(np.max(values)),
        f"{prefix}_iqr": float(p75 - p25),
        f"{prefix}_range": float(np.max(values) - np.min(values)),
        f"{prefix}_mad": float(np.mean(np.abs(values - mean_value))),
        f"{prefix}_rmad": float(np.mean(np.abs(robust - robust_mean))),
        f"{prefix}_medad": float(np.median(np.abs(values - median_value))),
        f"{prefix}_cov": float(std_value / mean_value) if abs(mean_value) > 1e-12 else float("nan"),
        f"{prefix}_qcod": float((p75 - p25) / (p75 + p25)) if abs(p75 + p25) > 1e-12 else float("nan"),
    }
    return features


def _mask_values(image: sitk.Image, mask: sitk.Image) -> np.ndarray:
    image_arr = sitk.GetArrayFromImage(image).astype(np.float64, copy=False)
    mask_arr = sitk.GetArrayFromImage(mask).astype(bool, copy=False)
    return image_arr[mask_arr]


def _mask_bbox(mask: sitk.Image) -> tuple[int, int, int]:
    stats = sitk.LabelShapeStatisticsImageFilter()
    stats.Execute(mask)
    if not stats.HasLabel(IBSI1_LABEL):
        return (0, 0, 0)
    bbox = stats.GetBoundingBox(IBSI1_LABEL)
    return (int(bbox[3]), int(bbox[4]), int(bbox[5]))


def _mask_voxel_count(mask: sitk.Image) -> int:
    stats = sitk.LabelShapeStatisticsImageFilter()
    stats.Execute(mask)
    if not stats.HasLabel(IBSI1_LABEL):
        return 0
    return int(stats.GetNumberOfPixels(IBSI1_LABEL))


def _mask_center(mask: sitk.Image, values: np.ndarray | None = None) -> np.ndarray | None:
    mask_arr = sitk.GetArrayFromImage(mask).astype(bool, copy=False)
    coords = np.argwhere(mask_arr)
    if coords.size == 0:
        return None

    spacing = np.array(mask.GetSpacing()[::-1], dtype=np.float64)
    origin = np.array(mask.GetOrigin()[::-1], dtype=np.float64)
    direction = np.array(mask.GetDirection(), dtype=np.float64).reshape(mask.GetDimension(), mask.GetDimension())[::-1, ::-1]
    physical = origin + (coords * spacing) @ direction.T

    if values is None:
        return np.mean(physical, axis=0)

    weights = values.astype(np.float64, copy=False)
    weights = weights - float(np.min(weights)) + 1.0
    weight_sum = float(np.sum(weights))
    if abs(weight_sum) <= 1e-12:
        return None
    return np.sum(physical * weights[:, None], axis=0) / weight_sum


def _build_discretized_values(image: sitk.Image, mask: sitk.Image, discretization: dict) -> np.ndarray:
    from flash_radiomics.imageoperations import binImage

    image_arr = sitk.GetArrayFromImage(image).astype(np.float64, copy=False)
    mask_arr = sitk.GetArrayFromImage(mask).astype(bool, copy=False)
    if not np.any(mask_arr):
        return np.array([], dtype=np.float64)

    texture = dict(discretization.get("texture_and_ih", {}))
    kwargs: dict[str, float | int] = {}
    if texture.get("method") == "FBN":
        kwargs["binCount"] = int(texture["n_bins"])
    else:
        kwargs["binWidth"] = float(texture["bin_size_hu"])

    discretized, _ = binImage(image_arr, np.where(mask_arr), **kwargs)
    return discretized[mask_arr].astype(np.float64, copy=False)


def _manual_feature_values(prepared: PreparedConfigCase) -> dict[str, float]:
    result: dict[str, float] = {}

    reseg_values = _mask_values(prepared.interpolated_image, prepared.intensity_mask)

    def add_image_stats(prefix: str, image: sitk.Image) -> None:
        arr = sitk.GetArrayFromImage(image).astype(np.float64, copy=False)
        size = tuple(int(v) for v in image.GetSize())
        spacing = tuple(float(v) for v in image.GetSpacing())
        result[f"img_dim_x_{prefix}"] = float(size[0])
        result[f"img_dim_y_{prefix}"] = float(size[1])
        result[f"img_dim_z_{prefix}"] = float(size[2])
        result[f"vox_dim_x_{prefix}"] = float(spacing[0])
        result[f"vox_dim_y_{prefix}"] = float(spacing[1])
        result[f"vox_dim_z_{prefix}"] = float(spacing[2])
        result[f"mean_int_{prefix}"] = float(np.mean(arr))
        result[f"min_int_{prefix}"] = float(np.min(arr))
        result[f"max_int_{prefix}"] = float(np.max(arr))

    def add_roi_stats(prefix: str, image: sitk.Image, intensity_mask: sitk.Image, morph_mask: sitk.Image) -> None:
        mask_size = tuple(int(v) for v in intensity_mask.GetSize())
        result[f"int_mask_dim_x_{prefix}"] = float(mask_size[0])
        result[f"int_mask_dim_y_{prefix}"] = float(mask_size[1])
        result[f"int_mask_dim_z_{prefix}"] = float(mask_size[2])

        int_bbox = _mask_bbox(intensity_mask)
        morph_bbox = _mask_bbox(morph_mask)
        result[f"int_mask_bb_dim_x_{prefix}"] = float(int_bbox[0])
        result[f"int_mask_bb_dim_y_{prefix}"] = float(int_bbox[1])
        result[f"int_mask_bb_dim_z_{prefix}"] = float(int_bbox[2])
        result[f"morph_mask_bb_dim_x_{prefix}"] = float(morph_bbox[0])
        result[f"morph_mask_bb_dim_y_{prefix}"] = float(morph_bbox[1])
        result[f"morph_mask_bb_dim_z_{prefix}"] = float(morph_bbox[2])

        result[f"int_mask_vox_count_{prefix}"] = float(_mask_voxel_count(intensity_mask))
        result[f"morph_mask_vox_count_{prefix}"] = float(_mask_voxel_count(morph_mask))

        roi_values = _mask_values(image, intensity_mask)
        if roi_values.size > 0:
            result[f"int_mask_mean_int_{prefix}"] = float(np.mean(roi_values))
            result[f"int_mask_min_int_{prefix}"] = float(np.min(roi_values))
            result[f"int_mask_max_int_{prefix}"] = float(np.max(roi_values))

    add_image_stats("init_img", prepared.initial_image)
    add_image_stats("interp_img", prepared.interpolated_image)
    add_roi_stats("init_roi", prepared.initial_image, prepared.initial_mask, prepared.initial_mask)
    add_roi_stats("interp_roi", prepared.interpolated_image, prepared.interpolated_mask, prepared.interpolated_mask)
    add_roi_stats("reseg_roi", prepared.interpolated_image, prepared.intensity_mask, prepared.morph_mask)

    if reseg_values.size > 0:
        result.update(_basic_distribution_features(reseg_values, "stat"))
        result["stat_energy"] = float(np.sum(reseg_values**2))
        result["stat_rms"] = float(np.sqrt(np.mean(reseg_values**2)))

    morph_arr = sitk.GetArrayFromImage(prepared.morph_mask).astype(bool, copy=False)
    int_arr = sitk.GetArrayFromImage(prepared.intensity_mask).astype(bool, copy=False)
    img_arr = sitk.GetArrayFromImage(prepared.interpolated_image).astype(np.float64, copy=False)
    if np.any(morph_arr) and np.any(int_arr):
        morph_coords = np.argwhere(morph_arr).astype(np.float64, copy=False)
        int_coords = np.argwhere(int_arr).astype(np.float64, copy=False)
        int_values = img_arr[int_arr]
        int_sum = float(np.sum(int_values))
        if not np.isclose(int_sum, 0.0):
            com_morph = np.mean(morph_coords, axis=0)
            com_int = np.sum(int_coords * int_values[:, None], axis=0) / int_sum
            spacing_zyx = np.asarray(prepared.interpolated_image.GetSpacing()[::-1], dtype=np.float64)
            result["morph_com"] = float(np.sqrt(np.sum(((com_morph - com_int) * spacing_zyx) ** 2.0)))

    discretized_values = _build_discretized_values(prepared.interpolated_image, prepared.intensity_mask, prepared.discretization)
    if discretized_values.size > 0:
        result.update(_basic_distribution_features(discretized_values, "ih"))
        hist_counts = np.unique(discretized_values.astype(np.int64, copy=False), return_counts=True)[1].astype(np.float64)
        probs = hist_counts / float(np.sum(hist_counts))
        result["ih_mode"] = _mode_from_values(discretized_values)
        result["ih_entropy"] = float(-np.sum(probs * np.log2(probs + np.finfo(np.float64).eps)))
        result["ih_uniformity"] = float(np.sum(probs**2))

    result["ngl_dc_perc_2D"] = 1.0
    result["ngl_dc_perc_2_5D"] = 1.0
    result["ngl_dc_perc_3D"] = 1.0
    result.pop("stat_medad", None)
    result.pop("ih_medad", None)
    return result


def _shape_derived_features(shape_results: dict[str, float], manual: dict[str, float]) -> dict[str, float]:
    mapped: dict[str, float] = {}
    for tag_map in (SHAPE_RESULT_TAGS, VALIDATION_SHAPE_EXTRA_TAGS):
        for feature_name, tag in tag_map.items():
            key = f"original_shape_{feature_name}"
            if key in shape_results:
                mapped[tag] = float(shape_results[key])

    volume = mapped.get("morph_volume")
    surface = mapped.get("morph_area_mesh")
    sphericity = mapped.get("morph_sphericity")
    if volume is not None and surface is not None and surface > 0.0:
        mapped.setdefault("morph_comp_1", float(volume / (surface ** 1.5 * math.sqrt(math.pi))))
        mapped.setdefault("morph_comp_2", float((36.0 * math.pi) * (volume**2.0) / (surface**3.0)))
        if sphericity is not None and sphericity > 0.0:
            mapped.setdefault("morph_sph_dispr", float(1.0 / sphericity))
            mapped.setdefault("morph_asphericity", float((1.0 / sphericity) - 1.0))
        if volume > 0.0:
            mapped["morph_integ_int"] = float(manual.get("stat_mean", float("nan")) * volume)
    if "morph_com" in manual:
        mapped["morph_com"] = float(manual["morph_com"])
    return mapped


def _compute_bounding_box_densities(shape_results: dict[str, float], prepared: PreparedConfigCase) -> dict[str, float]:
    mapped: dict[str, float] = {}
    volume = shape_results.get("original_shape_MeshVolume")
    area = shape_results.get("original_shape_SurfaceArea")
    if volume is None or area is None:
        return mapped

    bbox = _mask_bbox(prepared.morph_mask)
    spacing = np.array(prepared.morph_mask.GetSpacing(), dtype=np.float64)
    bbox_physical = np.array([bbox[0] * spacing[0], bbox[1] * spacing[1], bbox[2] * spacing[2]], dtype=np.float64)

    if np.all(bbox_physical > 0):
        bbox_volume = float(np.prod(bbox_physical))
        bbox_area = float(2.0 * (bbox_physical[0] * bbox_physical[1] + bbox_physical[0] * bbox_physical[2] + bbox_physical[1] * bbox_physical[2]))
        mapped["morph_vol_dens_aabb"] = float(volume / bbox_volume)
        mapped["morph_area_dens_aabb"] = float(area / bbox_area)

    major = shape_results.get("original_shape_MajorAxisLength")
    minor = shape_results.get("original_shape_MinorAxisLength")
    least = shape_results.get("original_shape_LeastAxisLength")
    if major is not None and minor is not None and least is not None and major > 0.0 and minor > 0.0 and least > 0.0:
        a = float(major) / 2.0
        b = float(minor) / 2.0
        c = float(least) / 2.0
        ellipsoid_volume = float((4.0 / 3.0) * math.pi * a * b * c)
        if ellipsoid_volume > 0.0:
            mapped["morph_vol_dens_aee"] = float(volume / ellipsoid_volume)
    return mapped


def _texture_family_note(extraction_mode: str, feature_class: str) -> str:
    suffix = TEXTURE_SUFFIX_BY_MODE[extraction_mode][feature_class]
    return f"ibsi_family_assumption:{feature_class}:{suffix}"


def _coerce_numeric_scalar(value: object) -> float | None:
    if isinstance(value, (int, float, np.integer, np.floating)):
        return float(value)
    if (
        isinstance(value, np.ndarray)
        and value.size == 1
        and np.issubdtype(value.dtype, np.number)
    ):
        return float(value.reshape(-1)[0])
    return None


def _tool_feature_values(prepared: PreparedConfigCase, result_dict: dict[str, float]) -> tuple[dict[str, float], dict[str, str]]:
    mapped: dict[str, float] = {}
    notes: dict[str, str] = {}

    for key, value in result_dict.items():
        if not key.startswith("original_"):
            continue
        scalar = _coerce_numeric_scalar(value)
        if scalar is None:
            continue
        scalar_value = scalar
        for feature_name, tag in (
            dict(FIRSTORDER_TAGS) | dict(VALIDATION_FIRSTORDER_EXTRA_TAGS)
        ).items():
            if key == f"original_firstorder_{feature_name}":
                mapped[tag] = scalar_value
        for feature_name, tag_prefix in GLCM_TAGS.items():
            if key == f"original_glcm_{feature_name}":
                tag = f"{tag_prefix}_{TEXTURE_SUFFIX_BY_MODE[prepared.extraction_mode]['glcm']}"
                mapped[tag] = scalar_value
                notes[tag] = _texture_family_note(prepared.extraction_mode, "glcm")
        for feature_name, tag_prefix in GLRLM_TAGS.items():
            if key == f"original_glrlm_{feature_name}":
                tag = f"{tag_prefix}_{TEXTURE_SUFFIX_BY_MODE[prepared.extraction_mode]['glrlm']}"
                mapped[tag] = scalar_value
                notes[tag] = _texture_family_note(prepared.extraction_mode, "glrlm")
        for feature_name, tag_prefix in GLSZM_TAGS.items():
            if key == f"original_glszm_{feature_name}":
                tag = f"{tag_prefix}_{TEXTURE_SUFFIX_BY_MODE[prepared.extraction_mode]['glszm']}"
                mapped[tag] = scalar_value
                notes[tag] = _texture_family_note(prepared.extraction_mode, "glszm")
        for feature_name, tag_prefix in GLDZM_TAGS.items():
            if key == f"original_gldzm_{feature_name}":
                tag = f"{tag_prefix}_{TEXTURE_SUFFIX_BY_MODE[prepared.extraction_mode]['gldzm']}"
                mapped[tag] = scalar_value
                notes[tag] = _texture_family_note(prepared.extraction_mode, "gldzm")
        for feature_name, tag_prefix in (
            dict(GLDM_TAGS) | dict(VALIDATION_GLDM_EXTRA_TAGS)
        ).items():
            if key == f"original_gldm_{feature_name}":
                tag = f"{tag_prefix}_{TEXTURE_SUFFIX_BY_MODE[prepared.extraction_mode]['gldm']}"
                mapped[tag] = scalar_value
                notes[tag] = _texture_family_note(prepared.extraction_mode, "gldm")
        for feature_name, tag_prefix in NGTDM_TAGS.items():
            if key == f"original_ngtdm_{feature_name}":
                tag = f"{tag_prefix}_{TEXTURE_SUFFIX_BY_MODE[prepared.extraction_mode]['ngtdm']}"
                mapped[tag] = scalar_value
                notes[tag] = _texture_family_note(prepared.extraction_mode, "ngtdm")
        for feature_name, tag in IH_TAGS.items():
            if key == f"original_ih_{feature_name}":
                mapped[tag] = scalar_value
        for feature_name, tag in IVH_TAGS.items():
            if key == f"original_ivh_{feature_name}":
                mapped[tag] = scalar_value

    return mapped, notes


def _normalize_config_id(value: str) -> str:
    text = str(value).strip()
    if text.lower().startswith("configuration "):
        return text.split()[-1].upper()
    return text.upper()


def load_ibsi1_configs(config_json_path: Path) -> dict[str, dict]:
    with config_json_path.open("r", encoding="utf-8") as handle:
        payload = json.load(handle)
    return {str(key).upper(): dict(value) for key, value in dict(payload.get("configs", {})).items()}


def load_ibsi1_reference_rows(template_path: Path) -> dict[str, list[dict]]:
    workbook = load_workbook(template_path, data_only=True)
    rows_by_config: dict[str, list[dict]] = {}
    for sheet in workbook.worksheets:
        config_id = _normalize_config_id(sheet.title.split()[-1])
        sheet_rows: list[dict] = []
        for values in sheet.iter_rows(min_row=2, values_only=True):
            dataset, family, feature, consensus, reference_value, tolerance, *_rest, tag = values
            if not family or not feature or not tag:
                continue
            sheet_rows.append(
                {
                    "config_id": config_id,
                    "dataset": str(dataset or ""),
                    "family": str(family),
                    "feature": str(feature),
                    "consensus": str(consensus or ""),
                    "reference_value": None if reference_value is None else float(reference_value),
                    "tolerance": None if tolerance is None else float(tolerance),
                    "tag": str(tag),
                }
            )
        rows_by_config[config_id] = sheet_rows
    return rows_by_config


def _tag_family(tag: str) -> str:
    text = str(tag)
    if text.startswith("morph_"):
        return "morphology"
    if text.startswith("stat_"):
        return "statistics"
    if text.startswith("ih_"):
        return "intensity_histogram"
    if text.startswith("ivh_"):
        return "intensity_volume_histogram"
    if text.startswith("cm_"):
        return "glcm"
    if text.startswith("rlm_"):
        return "glrlm"
    if text.startswith("szm_"):
        return "glszm"
    if text.startswith("dzm_"):
        return "gldzm"
    if text.startswith("ngt_"):
        return "ngtdm"
    if text.startswith("ngl_"):
        return "ngldm"
    return ""


def load_high_consensus_feature_tags(map_csv_path: Path = DEFAULT_IBSI_HIGH_CONSENSUS_MAP) -> set[str]:
    with map_csv_path.open("r", newline="", encoding="utf-8") as handle:
        reader = csv.DictReader(handle)
        tags = {str(row.get("mirp", "")).strip() for row in reader}
    tags.discard("")
    if not tags:
        raise ValueError(f"No MIRP tags found in {map_csv_path}.")
    return tags


def load_tool_mapped_feature_tags(
    tool_name: str,
    map_csv_path: Path = DEFAULT_IBSI_HIGH_CONSENSUS_MAP,
) -> set[str]:
    tool_column = {
        "flash": "flash",
        "mirp": "mirp",
        "pyradiomics": "pyr",
    }.get(str(tool_name).strip().lower())
    if tool_column is None:
        return set()

    with map_csv_path.open("r", newline="", encoding="utf-8") as handle:
        reader = csv.DictReader(handle)
        tags = {
            str(row.get("mirp", "")).strip()
            for row in reader
            if str(row.get(tool_column, "")).strip()
        }
    tags.discard("")
    return tags


def _is_high_consensus_reference_tag(tag: str, allowed_base_tags: set[str]) -> bool:
    text = str(tag)
    base_tag = text
    for suffix in ("_2D_avg", "_2D_comb", "_2_5D_avg", "_2_5D_comb", "_2_5D", "_3D_avg", "_3D_comb", "_3D", "_2D"):
        if text.endswith(suffix):
            base_tag = text[: -len(suffix)]
            break
    return base_tag in allowed_base_tags


def _expected_segment_reference_tag(config_id: str, base_tag: str) -> str:
    family = _tag_family(base_tag)
    config = str(config_id).upper()
    if family in {"glcm", "glrlm"}:
        return f"{base_tag}_2_5D_avg" if config in {"A", "B"} else f"{base_tag}_3D_avg"
    if family in {"glszm", "gldzm", "ngtdm", "ngldm"}:
        return f"{base_tag}_2_5D" if config in {"A", "B"} else f"{base_tag}_3D"
    return base_tag


def _is_default_segment_high_consensus_reference_tag(
    config_id: str,
    tag: str,
    allowed_base_tags: set[str],
) -> bool:
    text = str(tag)
    for base_tag in allowed_base_tags:
        if text == _expected_segment_reference_tag(config_id, base_tag):
            return True
    return False


def filter_reference_rows_to_high_consensus(
    reference_rows_by_config: dict[str, list[dict]],
    allowed_base_tags: set[str],
    default_segment_mode_only: bool = True,
) -> tuple[dict[str, list[dict]], dict[str, dict]]:
    filtered: dict[str, list[dict]] = {}
    filter_stats: dict[str, dict] = {}
    for config_id, rows in reference_rows_by_config.items():
        if default_segment_mode_only and str(config_id).upper() in {"A", "B", "C", "D", "E"}:
            kept = [
                row
                for row in rows
                if _is_default_segment_high_consensus_reference_tag(
                    config_id,
                    str(row["tag"]),
                    allowed_base_tags,
                )
            ]
        else:
            kept = [row for row in rows if _is_high_consensus_reference_tag(str(row["tag"]), allowed_base_tags)]
        skipped = len(rows) - len(kept)
        filtered[config_id] = kept
        filter_stats[config_id] = {
            "reference_rows_input": int(len(rows)),
            "reference_rows_kept": int(len(kept)),
            "reference_rows_skipped": int(skipped),
            "default_segment_mode_only": bool(default_segment_mode_only),
        }
    return filtered, filter_stats


def _build_texture_settings(config: dict) -> tuple[dict, bool, int]:
    extraction_mode = str(config.get("extraction_mode", "3D")).upper()
    texture_settings = {
        "additionalInfo": False,
        "distances": [int(config.get("texture_parameters", {}).get("glcm_distance", 1))],
        "gldm_a": int(config.get("texture_parameters", {}).get("ngldm_coarseness_alpha", 0)),
        "force2D": bool(extraction_mode == "2D"),
        "force2Ddimension": 0,
    }

    discretization = dict(config.get("discretization", {}).get("texture_and_ih", {}))
    resegmentation_range = config.get("resegmentation", {}).get("range_hu")
    if discretization.get("method") == "FBN":
        texture_settings["binCount"] = int(discretization["n_bins"])
    else:
        texture_settings["binWidth"] = float(discretization["bin_size_hu"])
        if isinstance(resegmentation_range, list) and len(resegmentation_range) == 2:
            texture_settings["binMinimum"] = float(min(resegmentation_range))

    ivh_discretization = config.get("discretization", {}).get("ivh")
    if isinstance(ivh_discretization, dict):
        if ivh_discretization.get("method") == "FBN":
            texture_settings["ivhDiscretizationMethod"] = "fixed_bin_number"
            texture_settings["ivhBinCount"] = int(ivh_discretization["n_bins"])
        elif ivh_discretization.get("method") == "FBS":
            texture_settings["ivhDiscretizationMethod"] = "fixed_bin_size"
            texture_settings["ivhBinWidth"] = float(ivh_discretization["bin_size_hu"])
    else:
        texture_settings["ivhDiscretizationMethod"] = "none"

    if isinstance(resegmentation_range, list) and len(resegmentation_range) == 2:
        texture_settings["ivhIntensityRange"] = tuple(sorted(float(v) for v in resegmentation_range))

    return texture_settings, bool(extraction_mode == "2D"), 0


def prepare_ibsi1_case(config_id: str, config: dict, data_root: Path) -> PreparedConfigCase:
    image_path = data_root / "ibsi_1_ct_radiomics_phantom" / "nifti" / "image" / "phantom.nii.gz"
    mask_path = data_root / "ibsi_1_ct_radiomics_phantom" / "nifti" / "mask" / "mask.nii.gz"
    if not image_path.exists() or not mask_path.exists():
        raise FileNotFoundError("IBSI-1 CT phantom NIfTI files were not found under reproducibility/data/ibsi/data_sets")

    initial_image, initial_mask = _read_rounded_phantom(image_path, mask_path)
    interpolated_image, interpolated_mask = _resample_case(initial_image, initial_mask, config)
    intensity_mask = _apply_resegmentation(interpolated_image, interpolated_mask, config)
    morph_mask = sitk.Image(interpolated_mask)

    texture_settings, force2d, force2d_dimension = _build_texture_settings(config)
    note_parts = [
        f"nifti_rounded_to_nearest_integer:{config_id}",
        f"texture_mode_assumption:{str(config.get('extraction_mode', '3D')).upper()}",
    ]
    return PreparedConfigCase(
        config_id=config_id,
        extraction_mode=str(config.get("extraction_mode", "3D")).upper(),
        note="|".join(note_parts),
        initial_image=initial_image,
        initial_mask=initial_mask,
        interpolated_image=interpolated_image,
        interpolated_mask=interpolated_mask,
        morph_mask=morph_mask,
        intensity_mask=intensity_mask,
        discretization=dict(config.get("discretization", {})),
        force2d=force2d,
        force2d_dimension=force2d_dimension,
        texture_settings=texture_settings,
    )


def _timing_stats(durations: list[float]) -> dict[str, float]:
    arr = np.array(durations, dtype=np.float64)
    if arr.size == 0:
        return {
            "mean_wall_time_s": float("nan"),
            "median_wall_time_s": float("nan"),
            "std_wall_time_s": float("nan"),
            "min_wall_time_s": float("nan"),
            "max_wall_time_s": float("nan"),
        }
    return {
        "mean_wall_time_s": float(np.mean(arr)),
        "median_wall_time_s": float(np.median(arr)),
        "std_wall_time_s": float(np.std(arr)),
        "min_wall_time_s": float(np.min(arr)),
        "max_wall_time_s": float(np.max(arr)),
    }


def _run_tool_once(
    prepared: PreparedConfigCase,
    *,
    build_extractor: Callable[[tuple[str, ...], dict], object],
) -> dict[str, float]:
    result: dict[str, float] = {}

    shape_extractor = build_extractor(("shape",), {"additionalInfo": False})
    shape_result = shape_extractor.execute(
        prepared.interpolated_image,
        prepared.morph_mask,
        label=IBSI1_LABEL,
    )
    result.update({str(key): value for key, value in dict(shape_result).items()})

    feature_classes = (
        "firstorder",
        "glcm",
        "glrlm",
        "glszm",
        "gldzm",
        "gldm",
        "ngtdm",
        "ih",
        "ivh",
    )
    feature_settings = dict(prepared.texture_settings)
    feature_settings["gldzmDistanceMask"] = prepared.morph_mask
    feature_extractor = build_extractor(feature_classes, feature_settings)
    feature_result = feature_extractor.execute(
        prepared.interpolated_image,
        prepared.intensity_mask,
        label=IBSI1_LABEL,
    )
    result.update({str(key): value for key, value in dict(feature_result).items()})
    return result


def _evaluate_validation_rows(
    reference_rows: list[dict],
    values_by_tag: dict[str, float],
    source_by_tag: dict[str, str],
    notes_by_tag: dict[str, str],
) -> list[dict]:
    evaluated_rows: list[dict] = []
    for row in reference_rows:
        tag = str(row["tag"])
        reference_value = row["reference_value"]
        tolerance = row["tolerance"]
        result_value = values_by_tag.get(tag)
        source_key = source_by_tag.get(tag, "")
        note = notes_by_tag.get(tag, "")

        status = "unsupported"
        difference = float("nan")
        passed = ""

        if reference_value is None or tolerance is None:
            status = "not_standardized"
        elif result_value is None:
            status = "unsupported"
        elif not math.isfinite(float(result_value)):
            status = "error"
        else:
            difference = abs(float(result_value) - float(reference_value))
            numeric_noise_tol = 1e-4
            within_tolerance = difference <= float(tolerance) or math.isclose(
                float(result_value),
                float(reference_value),
                rel_tol=0.0,
                abs_tol=numeric_noise_tol,
            )
            status = "pass" if within_tolerance else "fail"
            passed = "1" if status == "pass" else "0"

        evaluated_rows.append(
            {
                "config_id": row["config_id"],
                "dataset": row["dataset"],
                "family": row["family"],
                "feature": row["feature"],
                "consensus": row["consensus"],
                "reference_value": "" if reference_value is None else float(reference_value),
                "tolerance": "" if tolerance is None else float(tolerance),
                "result_value": "" if result_value is None else float(result_value),
                "absolute_difference": difference,
                "status": status,
                "passed": passed,
                "source_key": source_key,
                "note": note,
                "tag": tag,
            }
        )
    return evaluated_rows


def run_ibsi1_segment_validation(
    *,
    tool_name: str,
    backend: str,
    exp_dir: Path,
    repeats: int,
    data_root: Path,
    template_path: Path,
    config_json_path: Path,
    build_extractor: Callable[[tuple[str, ...], dict], object],
    high_consensus_only: bool = True,
    high_consensus_map_path: Path = DEFAULT_IBSI_HIGH_CONSENSUS_MAP,
) -> None:
    configs = load_ibsi1_configs(config_json_path)
    reference_rows_by_config = load_ibsi1_reference_rows(template_path)
    reference_filter_stats: dict[str, dict] = {}
    high_consensus_tags: set[str] = set()
    tool_mapped_tags: set[str] = set()
    if high_consensus_only:
        high_consensus_tags = load_high_consensus_feature_tags(high_consensus_map_path)
        tool_mapped_tags = load_tool_mapped_feature_tags(tool_name, high_consensus_map_path)
        reference_rows_by_config, reference_filter_stats = filter_reference_rows_to_high_consensus(
            reference_rows_by_config,
            high_consensus_tags,
        )

    validation_rows: list[dict] = []
    timing_rows: list[dict] = []
    summary_rows: list[dict] = []

    for config_id in sorted(configs):
        prepared = prepare_ibsi1_case(config_id, configs[config_id], data_root)

        durations: list[float] = []
        last_result: dict[str, float] = {}
        for _ in range(max(1, int(repeats))):
            start = time.perf_counter()
            last_result = _run_tool_once(prepared, build_extractor=build_extractor)
            durations.append(float(time.perf_counter() - start))

        manual_values = _manual_feature_values(prepared)
        tool_values, tool_notes = _tool_feature_values(prepared, last_result)

        shape_values: dict[str, float] = {}
        for key, value in last_result.items():
            scalar = _coerce_numeric_scalar(value)
            if str(key).startswith("original_shape_") and scalar is not None:
                shape_values[str(key)] = scalar
        derived_shape = _shape_derived_features(shape_values, manual_values)
        for tag, value in _compute_bounding_box_densities(shape_values, prepared).items():
            derived_shape.setdefault(tag, value)

        values_by_tag = dict(manual_values)
        values_by_tag.update(tool_values)
        values_by_tag.update(derived_shape)

        if tool_mapped_tags:
            values_by_tag = {
                tag: value
                for tag, value in values_by_tag.items()
                if _is_high_consensus_reference_tag(tag, tool_mapped_tags)
            }

        source_by_tag = {tag: tag for tag in manual_values}
        source_by_tag.update({tag: f"tool:{tool_name}" for tag in tool_values})
        source_by_tag.update({tag: "shape_derived" for tag in derived_shape})

        if tool_mapped_tags:
            source_by_tag = {
                tag: source
                for tag, source in source_by_tag.items()
                if tag in values_by_tag
            }

        notes_by_tag = {tag: prepared.note for tag in manual_values}
        notes_by_tag.update(tool_notes)
        for tag in derived_shape:
            notes_by_tag.setdefault(tag, "derived_from_shape_outputs")
        if tool_mapped_tags:
            notes_by_tag = {
                tag: note
                for tag, note in notes_by_tag.items()
                if tag in values_by_tag
            }

        config_rows = _evaluate_validation_rows(reference_rows_by_config.get(config_id, []), values_by_tag, source_by_tag, notes_by_tag)
        validation_rows.extend(config_rows)

        status_counts = {status: 0 for status in ALL_ROWS_STATUS_ORDER}
        for row in config_rows:
            status_counts[str(row["status"])] += 1

        comparable_rows = status_counts["pass"] + status_counts["fail"]
        timing = _timing_stats(durations)
        timing_rows.append(
            {
                "config_id": config_id,
                "case_id": IBSI1_CASE_ID,
                "label": IBSI1_LABEL,
                "roi_name": IBSI1_ROI_NAME,
                "tool": tool_name,
                "backend": backend,
                "mode": "segment",
                "mean_wall_time_s": timing["mean_wall_time_s"],
                "median_wall_time_s": timing["median_wall_time_s"],
                "std_wall_time_s": timing["std_wall_time_s"],
                "min_wall_time_s": timing["min_wall_time_s"],
                "max_wall_time_s": timing["max_wall_time_s"],
                "feature_count": int(len(values_by_tag)),
                "comparable_rows": int(comparable_rows),
                "pass_count": int(status_counts["pass"]),
                "fail_count": int(status_counts["fail"]),
                "unsupported_count": int(status_counts["unsupported"]),
                "not_standardized_count": int(status_counts["not_standardized"]),
                "error_count": int(status_counts["error"]),
                "note": prepared.note,
            }
        )

        summary_rows.append(
            {
                "config_id": config_id,
                "extraction_mode": prepared.extraction_mode,
                "rows_total": int(len(config_rows)),
                "rows_comparable": int(comparable_rows),
                "rows_pass": int(status_counts["pass"]),
                "rows_fail": int(status_counts["fail"]),
                "rows_unsupported": int(status_counts["unsupported"]),
                "rows_not_standardized": int(status_counts["not_standardized"]),
                "rows_error": int(status_counts["error"]),
                "pass_rate": float(status_counts["pass"] / comparable_rows) if comparable_rows > 0 else float("nan"),
            }
        )

    validation_csv = exp_dir / f"{tool_name}_ibsi1_validation.csv"
    timing_csv = exp_dir / f"{tool_name}_ibsi1_case_timing.csv"
    summary_json = exp_dir / f"{tool_name}_ibsi1_validation_summary.json"

    validation_fieldnames = [
        "config_id",
        "dataset",
        "family",
        "feature",
        "consensus",
        "reference_value",
        "tolerance",
        "result_value",
        "absolute_difference",
        "status",
        "passed",
        "source_key",
        "note",
        "tag",
    ]
    timing_fieldnames = [
        "config_id",
        "case_id",
        "label",
        "roi_name",
        "tool",
        "backend",
        "mode",
        "mean_wall_time_s",
        "median_wall_time_s",
        "std_wall_time_s",
        "min_wall_time_s",
        "max_wall_time_s",
        "feature_count",
        "comparable_rows",
        "pass_count",
        "fail_count",
        "unsupported_count",
        "not_standardized_count",
        "error_count",
        "note",
    ]

    total_status = {status: 0 for status in ALL_ROWS_STATUS_ORDER}
    for row in validation_rows:
        total_status[str(row["status"])] += 1
    total_comparable = total_status["pass"] + total_status["fail"]

    write_csv(validation_csv, validation_fieldnames, validation_rows)
    write_csv(timing_csv, timing_fieldnames, timing_rows)
    write_json(
        summary_json,
        {
            "tool": tool_name,
            "backend": backend,
            "mode": "segment",
            "suite": "ibsi1",
            "repeat": int(repeats),
            "number_of_configs": int(len(summary_rows)),
            "config_ids": [str(row["config_id"]) for row in summary_rows],
            "high_consensus_only": bool(high_consensus_only),
            "high_consensus_feature_count": int(len(high_consensus_tags)) if high_consensus_only else 0,
            "tool_mapped_feature_count": int(len(tool_mapped_tags)) if high_consensus_only else 0,
            "high_consensus_map_csv": relative_to_workspace(high_consensus_map_path) if high_consensus_only else "",
            "reference_filter": reference_filter_stats,
            "rows_total": int(len(validation_rows)),
            "rows_comparable": int(total_comparable),
            "rows_pass": int(total_status["pass"]),
            "rows_fail": int(total_status["fail"]),
            "rows_unsupported": int(total_status["unsupported"]),
            "rows_not_standardized": int(total_status["not_standardized"]),
            "rows_error": int(total_status["error"]),
            "pass_rate": float(total_status["pass"] / total_comparable) if total_comparable > 0 else float("nan"),
            "validation_csv": relative_to_workspace(validation_csv),
            "timing_csv": relative_to_workspace(timing_csv),
            "configs": summary_rows,
            "notes": [
                "IBSI-1 CT phantom NIfTI image was rounded to nearest integer before processing.",
                "Texture-family mapping follows tool-style family assumptions encoded in ibsi1_validation_common.py.",
                "Only high-consensus rows from reproducibility/ibsi/ibsi_high_consensus_165_feature_map.csv are scored by default.",
                "Tool support is limited to rows with a non-empty tool-specific mapping in the same crosswalk.",
                "Rows without an official reference value or tolerance are reported as not_standardized.",
            ],
        },
    )

    print(f"exp_dir={relative_to_workspace(exp_dir)}")
    print(f"summary_json={relative_to_workspace(summary_json)}")
